#!/usr/bin/env python3
"""
Generate a complete Word documentation file for the AWS-LUMI project.

Output:
  docs/Transcribe_MVP_Complete_Documentation.docx
"""

from __future__ import annotations

import datetime as dt
import json
import re
import subprocess
from collections import Counter
from pathlib import Path
from typing import Dict, List, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


ROOT = Path(__file__).resolve().parents[1]
README = ROOT / "README.md"
TEMPLATE = ROOT / "infra" / "template.yaml"
APP_PY = ROOT / "backend" / "src" / "app.py"
ASR_WORKER_PY = ROOT / "backend" / "asr_worker" / "app.py"
SPEAKER_ID_PY = ROOT / "backend" / "speaker_id" / "app.py"
FRONTEND_REACT = ROOT / "frontend" / "src" / "App.jsx"
FRONTEND_STATIC = ROOT / "frontend" / "publish" / "index.html"
QA_TEST = ROOT / "qa" / "tests" / "e2e.spec.mjs"

DOCS_DIR = ROOT / "docs"
DOC_PATH = DOCS_DIR / "Transcribe_MVP_Complete_Documentation.docx"


def run_cmd(args: List[str], timeout: int = 60) -> Tuple[int, str, str]:
    p = subprocess.run(args, capture_output=True, text=True, timeout=timeout, cwd=str(ROOT))
    return p.returncode, p.stdout, p.stderr


def read_text(path: Path) -> str:
    return path.read_text(encoding="utf-8")


def latest_qa_run() -> Tuple[str, Path] | Tuple[None, None]:
    qa_runs = ROOT / "qa_runs"
    if not qa_runs.exists():
        return None, None
    runs = [d for d in qa_runs.iterdir() if d.is_dir() and d.name != "_wip" and (d / "report.md").exists()]
    if not runs:
        return None, None
    runs.sort(key=lambda d: d.stat().st_mtime, reverse=True)
    return runs[0].name, runs[0] / "report.md"


def get_stack_outputs(stack_name: str = "transcribe-mvp", region: str = "us-west-1") -> Dict[str, str]:
    code, out, _err = run_cmd(
        [
            "aws",
            "cloudformation",
            "describe-stacks",
            "--stack-name",
            stack_name,
            "--region",
            region,
            "--query",
            "Stacks[0].Outputs",
            "--output",
            "json",
        ],
        timeout=120,
    )
    if code != 0:
        return {}
    try:
        rows = json.loads(out)
    except json.JSONDecodeError:
        return {}
    result: Dict[str, str] = {}
    for row in rows:
        k = row.get("OutputKey")
        v = row.get("OutputValue")
        if k and isinstance(v, str):
            result[k] = v
    return result


def parse_routes(app_text: str) -> List[Tuple[str, str]]:
    routes: List[Tuple[str, str]] = []
    for line in app_text.splitlines():
        line = line.strip()
        m1 = re.search(r"if method == '([A-Z]+)' and path == '([^']+)'", line)
        if m1:
            routes.append((m1.group(1), m1.group(2)))
            continue
        m2 = re.search(r"if method == '([A-Z]+)' and re\.fullmatch\(r'([^']+)'", line)
        if m2:
            method = m2.group(1)
            raw = m2.group(2)
            normalized = (
                raw.replace(r"\/", "/")
                .replace("[^/]+", "{id}")
                .replace("{id}/status", "{transcriptId}/status")
                .replace("{id}/interactions", "{transcriptId}/interactions")
            )
            if normalized == "/transcriptions/{id}":
                normalized = "/transcriptions/{transcriptId}"
            routes.append((method, normalized))
    dedup: List[Tuple[str, str]] = []
    seen = set()
    for r in routes:
        if r not in seen:
            dedup.append(r)
            seen.add(r)
    return dedup


def parse_aws_resource_inventory(template_text: str) -> Dict[str, int]:
    types = re.findall(r"Type:\s*(AWS::[A-Za-z0-9:]+)", template_text)
    by_service = Counter()
    for t in types:
        parts = t.split("::")
        service = parts[1] if len(parts) >= 3 else t
        by_service[service] += 1
    return dict(sorted(by_service.items(), key=lambda kv: kv[0]))


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    r = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r'TOC \o "1-3" \h \z \u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    fld_text = OxmlElement("w:t")
    fld_text.text = "Right-click and update field to refresh the Table of Contents."
    fld_sep.append(fld_text)
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    r._r.append(fld_begin)
    r._r.append(instr)
    r._r.append(fld_sep)
    r._r.append(fld_end)


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_bullets(doc: Document, lines: List[str]) -> None:
    for line in lines:
        doc.add_paragraph(line, style="List Bullet")


def add_numbered(doc: Document, lines: List[str]) -> None:
    for line in lines:
        doc.add_paragraph(line, style="List Number")


def add_table_from_pairs(doc: Document, title: str, rows: List[Tuple[str, str]]) -> None:
    add_heading(doc, title, level=3)
    table = doc.add_table(rows=1, cols=2)
    hdr = table.rows[0].cells
    hdr[0].text = "Name"
    hdr[1].text = "Value"
    for k, v in rows:
        cells = table.add_row().cells
        cells[0].text = str(k)
        cells[1].text = str(v)


def add_code_block(doc: Document, code: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(code)
    run.font.name = "Consolas"
    run.font.size = Pt(10)


def summarize_qa_report(report_path: Path) -> List[str]:
    if not report_path or not report_path.exists():
        return ["No QA report file was found in qa_runs/."]
    text = read_text(report_path)
    out: List[str] = []
    for line in text.splitlines():
        if line.startswith("- "):
            out.append(line[2:].strip())
        if len(out) >= 10:
            break
    if not out:
        out = ["QA report exists but no bullet summary lines were parsed."]
    return out


def main() -> None:
    app_text = read_text(APP_PY)
    asr_text = read_text(ASR_WORKER_PY)
    speaker_text = read_text(SPEAKER_ID_PY)
    template_text = read_text(TEMPLATE)
    readme_text = read_text(README)

    routes = parse_routes(app_text)
    resource_inventory = parse_aws_resource_inventory(template_text)
    stack_outputs = get_stack_outputs()
    qa_run_id, qa_report_path = latest_qa_run()
    qa_summary = summarize_qa_report(qa_report_path) if qa_report_path else ["No QA report available."]

    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()

    # Title page
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = title.add_run("AWS-LUMI Speaker Transcription MVP")
    r.bold = True
    r.font.size = Pt(24)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Complete Project Documentation")
    doc.add_paragraph("")
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run(f"Environment: transcribe-mvp (us-west-1)\nDate: {dt.date.today().isoformat()}\nAuthor: Codex")
    doc.add_page_break()

    # Document control
    add_heading(doc, "1. Title & Governance", level=1)
    add_heading(doc, "1.1 Document Control", level=2)
    t = doc.add_table(rows=1, cols=4)
    t.rows[0].cells[0].text = "Version"
    t.rows[0].cells[1].text = "Date"
    t.rows[0].cells[2].text = "Owner"
    t.rows[0].cells[3].text = "Notes"
    row = t.add_row().cells
    row[0].text = "1.0"
    row[1].text = dt.date.today().isoformat()
    row[2].text = "Project Engineering Team"
    row[3].text = "Initial comprehensive documentation release."

    add_heading(doc, "1.2 Audience and Reading Guide", level=2)
    add_bullets(
        doc,
        [
            "Primary audience: mixed (technical and business stakeholders).",
            "Use Sections 2-4 for product and architecture understanding.",
            "Use Sections 7-11 for API, operations, deployment, and troubleshooting.",
            "Appendices provide snapshots, references, and QA traceability.",
        ],
    )

    add_heading(doc, "1.3 Table of Contents", level=2)
    add_toc(doc)

    # Executive summary
    add_heading(doc, "2. Executive Summary", level=1)
    doc.add_paragraph(
        "The AWS-LUMI project is an MVP web application that transcribes uploaded audio, applies diarization "
        "(who spoke when), and performs calibration-based speaker naming (Kid, Parent 1, Parent 2). "
        "It is designed for low idle cost by running ASR as on-demand ECS tasks and persisting transcript "
        "results in DynamoDB with optional SQL analytics in Aurora PostgreSQL."
    )
    add_bullets(
        doc,
        [
            "Current maturity: MVP in active iteration with automated E2E QA coverage.",
            "Security posture: HTTPS frontend via CloudFront + private S3 origin (OAC).",
            "Cost posture: scale-to-zero compute model with worker tasks started per transcription.",
            "No authentication framework (e.g., Cognito) is currently enabled in the MVP path.",
        ],
    )

    # Business + functional scope
    add_heading(doc, "3. Business + Functional Scope", level=1)
    add_heading(doc, "3.1 End-User Journey", level=2)
    add_numbered(
        doc,
        [
            "Enter user ID and click Login.",
            "Upload calibration audio for Kid/Parent 1/Parent 2.",
            "Upload or record conversation audio.",
            "Click Upload + Transcribe and monitor status messages.",
            "Review summary metrics, speaker labels, and full transcript.",
            "Review Kid progression chart by day/week/month.",
        ],
    )
    add_heading(doc, "3.2 In Scope", level=2)
    add_bullets(
        doc,
        [
            "Calibration-based speaker naming.",
            "ASR + diarization + speaker identification pipeline.",
            "Transcript persistence and analytics endpoints.",
            "Operational deploy scripts and QA artifacts.",
        ],
    )
    add_heading(doc, "3.3 Out of Scope", level=2)
    add_bullets(
        doc,
        [
            "User/password authentication and role-based authorization.",
            "Multi-tenant enterprise controls.",
            "Mobile-native application clients.",
        ],
    )

    # Architecture
    add_heading(doc, "4. Solution Architecture", level=1)
    add_heading(doc, "4.1 Component Map", level=2)
    add_bullets(
        doc,
        [
            "CloudFront + private S3: frontend hosting over HTTPS.",
            "HTTP API (API Gateway) + Lambda: orchestration and API contract.",
            "ECS ASR worker: Whisper speech-to-text + pyannote diarization.",
            "Speaker ID Lambda: SpeechBrain embedding matching vs calibrations.",
            "DynamoDB: transcript/session persistence + stage/status state.",
            "Aurora PostgreSQL (optional): analytics and interaction queries.",
        ],
    )
    add_heading(doc, "4.2 Processing Sequence", level=2)
    add_numbered(
        doc,
        [
            "Cold start: start worker task.",
            "Speech-to-text: transcribe audio with Whisper.",
            "Diarization: identify turn boundaries by speaker.",
            "Speaker identification: apply calibration matching.",
            "Finalize transcript and persist outputs.",
            "Expose status and results via API polling.",
        ],
    )
    add_heading(doc, "4.3 Stage Message Model", level=2)
    add_bullets(
        doc,
        [
            "Cold start: starting worker...",
            "Speech-to-text: transcribing audio...",
            "Diarization: detecting who spoke when...",
            "Speaker identification: matching Kid/Parent calibrations...",
            "Finalizing transcript...",
            "Transcription complete.",
            "Failed: <reason>",
        ],
    )

    # Infrastructure
    add_heading(doc, "5. Infrastructure (AWS)", level=1)
    add_heading(doc, "5.1 Resource Inventory by Service", level=2)
    inv_rows = [(svc, str(count)) for svc, count in resource_inventory.items()]
    add_table_from_pairs(doc, "Inventory Count (from infra/template.yaml)", inv_rows)

    add_heading(doc, "5.2 Current Deployed Stack Outputs", level=2)
    if stack_outputs:
        output_rows = sorted(stack_outputs.items(), key=lambda kv: kv[0])
        add_table_from_pairs(doc, "CloudFormation Outputs (transcribe-mvp, us-west-1)", output_rows)
    else:
        doc.add_paragraph("Stack outputs could not be resolved at generation time.")

    add_heading(doc, "5.3 Security Posture", level=2)
    add_bullets(
        doc,
        [
            "Frontend bucket is private; read access is restricted through CloudFront OAC.",
            "Frontend delivery uses HTTPS endpoint through CloudFront.",
            "API remains on dedicated HTTP API endpoint with CORS handling in Lambda.",
            "No plaintext secrets are documented in this file.",
        ],
    )

    add_heading(doc, "5.4 Cost Model", level=2)
    add_bullets(
        doc,
        [
            "ASR workload executes as on-demand ECS tasks (scale-to-zero when idle).",
            "Cold start latency can occur when launching worker compute from idle.",
            "CPU path is the default runtime baseline; GPU is disabled unless explicitly enabled.",
        ],
    )

    # Data model & persistence
    add_heading(doc, "6. Data Model & Persistence", level=1)
    add_heading(doc, "6.1 DynamoDB (Transcripts) Core Fields", level=2)
    add_bullets(
        doc,
        [
            "userId, transcriptId, status, createdAt, updatedAt",
            "audioS3Key, transcriptJsonS3Key, numSpeakers",
            "segments[], speakerStats[], fullText",
            "processingStage, processingStageUpdatedAt",
            "effectiveDate, userTimeZone, logicalDay",
        ],
    )
    add_heading(doc, "6.2 Daily Word Stats Table", level=2)
    add_bullets(
        doc,
        [
            "Stores per-user/day/role word counts for Kid/Client1/Client2.",
            "Tracks uniqueWordCount and totalWordCount with token map.",
            "Used as fallback analytics source when SQL has no rows.",
        ],
    )
    add_heading(doc, "6.3 SQL Analytics (Aurora)", level=2)
    add_bullets(
        doc,
        [
            "Maintains transcript, interactions, metrics, and token occurrences.",
            "Uses logical_day semantics for user-time alignment (fallback to created_at date).",
            "Serves daily/weekly/monthly and latency analytics endpoints.",
        ],
    )
    add_heading(doc, "6.4 S3 Key Conventions", level=2)
    add_bullets(
        doc,
        [
            "uploads/<uuid>-<filename>",
            "transcripts/<transcriptId>/raw_whisper.json (and related artifacts)",
            "calibrations/<userId>/kid|parent1|parent2",
        ],
    )

    # API reference
    add_heading(doc, "7. API Reference", level=1)
    add_heading(doc, "7.1 Live Route Catalog (from backend/src/app.py)", level=2)
    route_table = doc.add_table(rows=1, cols=2)
    route_table.rows[0].cells[0].text = "Method"
    route_table.rows[0].cells[1].text = "Path"
    for method, path in routes:
        cells = route_table.add_row().cells
        cells[0].text = method
        cells[1].text = path

    add_heading(doc, "7.2 Status Lifecycle and Stage Semantics", level=2)
    add_bullets(
        doc,
        [
            "Transcript status values: QUEUED, PROCESSING, LABELING, COMPLETED, FAILED.",
            "Status endpoint returns external status: IN_PROGRESS/COMPLETED/FAILED.",
            "Optional stage field: COLD_START, ASR, DIARIZATION, SPEAKER_IDENTIFICATION, FINALIZING, COMPLETED, FAILED.",
        ],
    )

    add_heading(doc, "7.3 Example API Calls", level=2)
    add_code_block(
        doc,
        "POST /upload-url\n"
        'Body: {"fileName":"audio.mp3","contentType":"audio/mpeg"}\n\n'
        "POST /v1/calibration/presign\n"
        'Body: {"role":"kid","userId":"user1","fileName":"kid.mp3","contentType":"audio/mpeg"}\n\n'
        "POST /transcriptions\n"
        'Body: {"userId":"user1","s3Key":"uploads/...","effectiveDate":"YYYY-MM-DD","userTimeZone":"America/New_York"}\n\n'
        "GET /transcriptions/{transcriptId}/status\n"
        "GET /transcriptions/{transcriptId}\n"
    )

    # Frontend & UX
    add_heading(doc, "8. Frontend & UX Behavior", level=1)
    add_bullets(
        doc,
        [
            "UI supports login, calibration upload/recording, and upload+transcribe.",
            "Record Date (user time) can override logical day for writes and analytics reads.",
            "Kid progression chart supports day/week/month windows.",
            "Message field surfaces stage-aware status updates from backend polling.",
        ],
    )

    # Dev & deploy
    add_heading(doc, "9. Dev, Build, and Deploy Runbooks", level=1)
    add_heading(doc, "9.1 Full Deployment", level=2)
    add_code_block(doc, "powershell -ExecutionPolicy Bypass -File infra\\deploy.ps1 -StackName transcribe-mvp -Region us-west-1")
    add_heading(doc, "9.2 Fast Iteration Paths", level=2)
    add_code_block(
        doc,
        "API only:\n"
        "powershell -ExecutionPolicy Bypass -File infra\\deploy_api_only.ps1 -StackName transcribe-mvp -Region us-west-1\n\n"
        "Frontend only:\n"
        "powershell -ExecutionPolicy Bypass -File infra\\publish_frontend.ps1 -StackName transcribe-mvp -Region us-west-1\n\n"
        "ASR worker image:\n"
        "powershell -ExecutionPolicy Bypass -File infra\\build_asr_worker.ps1 -StackName transcribe-mvp -Region us-west-1\n\n"
        "Speaker-ID image:\n"
        "powershell -ExecutionPolicy Bypass -File infra\\build_speakerid.ps1 -StackName transcribe-mvp -Region us-west-1\n"
    )

    # QA
    add_heading(doc, "10. QA & Validation", level=1)
    if qa_run_id:
        doc.add_paragraph(f"Latest QA run: {qa_run_id}")
        doc.add_paragraph(f"Report: {qa_report_path.relative_to(ROOT)}")
    add_heading(doc, "10.1 Latest QA Summary", level=2)
    add_bullets(doc, qa_summary)
    add_heading(doc, "10.2 Acceptance Checklist", level=2)
    add_numbered(
        doc,
        [
            "Routes documented match live handler definitions.",
            "Architecture and resources reflect template + deployed outputs.",
            "Stage messages/lifecycle align with backend implementation.",
            "No secrets (tokens/passwords) included.",
        ],
    )

    # Operations & troubleshooting
    add_heading(doc, "11. Operations & Troubleshooting", level=1)
    add_heading(doc, "11.1 Common Incidents", level=2)
    add_bullets(
        doc,
        [
            "Slow transcription due to cold start.",
            "Calibration upload failures (signature/CORS/metadata mismatch).",
            "Speaker identification NO_MATCH outcomes.",
            "GPU capacity unavailability and fallback behavior.",
        ],
    )
    add_heading(doc, "11.2 Diagnostics", level=2)
    add_bullets(
        doc,
        [
            "Check API Lambda logs for route/status transitions.",
            "Check ASR worker logs for stage timings and diarization errors.",
            "Check speaker_id logs for embedding/similarity/match diagnostics.",
            "Use QA report network captures for request/response verification.",
        ],
    )

    # Roadmap
    add_heading(doc, "12. Roadmap / Next Improvements", level=1)
    add_numbered(
        doc,
        [
            "Add authentication and authorization controls.",
            "Reduce cold-start latency and improve stage visibility granularity.",
            "Expand analytics UX (calendar drill-downs, comparative trends).",
            "Harden operational monitoring, alerting, and rollback automation.",
        ],
    )

    # Appendices
    add_heading(doc, "13. Appendices", level=1)
    add_heading(doc, "13.1 Glossary", level=2)
    add_bullets(
        doc,
        [
            "ASR: Automatic Speech Recognition.",
            "Diarization: identifying who spoke when.",
            "OAC: Origin Access Control (CloudFront->S3 private access).",
            "Logical Day: user-time-aligned date key for analytics.",
        ],
    )
    add_heading(doc, "13.2 Resource Snapshot", level=2)
    if stack_outputs:
        add_table_from_pairs(doc, "Selected Key Outputs", sorted(stack_outputs.items(), key=lambda kv: kv[0]))
    else:
        doc.add_paragraph("Stack outputs were unavailable during generation.")

    add_heading(doc, "13.3 QA Artifact References", level=2)
    if qa_run_id:
        run_dir = ROOT / "qa_runs" / qa_run_id
        for path in sorted(run_dir.rglob("*")):
            if path.is_file():
                doc.add_paragraph(str(path.relative_to(ROOT)), style="List Bullet")
    else:
        doc.add_paragraph("No QA artifact directory found.")

    # Source traceability section
    add_heading(doc, "13.4 Source Traceability", level=2)
    add_bullets(
        doc,
        [
            f"README: {README.relative_to(ROOT)}",
            f"Infrastructure template: {TEMPLATE.relative_to(ROOT)}",
            f"API orchestrator: {APP_PY.relative_to(ROOT)}",
            f"ASR worker: {ASR_WORKER_PY.relative_to(ROOT)}",
            f"Speaker ID worker: {SPEAKER_ID_PY.relative_to(ROOT)}",
            f"Frontend (React): {FRONTEND_REACT.relative_to(ROOT)}",
            f"Frontend (static): {FRONTEND_STATIC.relative_to(ROOT)}",
            f"QA test suite: {QA_TEST.relative_to(ROOT)}",
        ],
    )

    doc.save(str(DOC_PATH))
    print(f"Generated: {DOC_PATH}")


if __name__ == "__main__":
    main()


